import os
import uuid
from datetime import datetime, timezone
from io import BytesIO

import boto3
import json
from docx import Document
from pypdf import PdfReader


s3_client = boto3.client("s3")

ALLOWED_CONTENT_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "text/plain"
}

PRESIGNED_URL_EXPIRATION = 900
MAX_PROCESSING_FILE_SIZE = 10 * 1024 * 1024  # 10 MB


def generate_upload_url(event: dict) -> dict:
    file_name = event.get("file_name")
    content_type = event.get("content_type")

    if not file_name or not content_type:
        return {
            "success": False,
            "error": "file_name y content_type son obligatorios"
        }

    if content_type not in ALLOWED_CONTENT_TYPES:
        return {
            "success": False,
            "error": "Tipo de archivo no permitido"
        }

    bucket_name = os.environ["DOCUMENTS_BUCKET"]

    now = datetime.now(timezone.utc)
    unique_id = str(uuid.uuid4())

    safe_file_name = os.path.basename(file_name)

    object_key = (
        f"uploads/{now.year}/"
        f"{now.month:02d}/"
        f"{unique_id}-{safe_file_name}"
    )

    upload_url = s3_client.generate_presigned_url(
        ClientMethod="put_object",
        Params={
            "Bucket": bucket_name,
            "Key": object_key,
            "ContentType": content_type
        },
        ExpiresIn=PRESIGNED_URL_EXPIRATION
    )

    return {
        "success": True,
        "action": "upload_url_generated",
        "bucket": bucket_name,
        "object_key": object_key,
        "upload_url": upload_url,
        "expires_in": PRESIGNED_URL_EXPIRATION
    }


def extract_txt(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode(
            "latin-1",
            errors="replace"
        )


def extract_pdf(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))

    pages = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            pages.append(text.strip())

    return "\n\n".join(pages)


def extract_docx(data: bytes) -> str:
    document = Document(BytesIO(data))

    blocks = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            blocks.append(text)

    for table in document.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]

            row_text = " | ".join(
                value
                for value in values
                if value
            )

            if row_text:
                blocks.append(row_text)

    return "\n".join(blocks)


def extract_document_text(event: dict) -> dict:
    object_key = event.get("object_key")

    if not object_key:
        return {
            "success": False,
            "error": "object_key es obligatorio"
        }

    bucket_name = os.environ["DOCUMENTS_BUCKET"]

    response = s3_client.get_object(
        Bucket=bucket_name,
        Key=object_key
    )

    content_length = response.get(
        "ContentLength",
        0
    )

    if content_length > MAX_PROCESSING_FILE_SIZE:
        return {
            "success": False,
            "error": (
                "El archivo supera el límite de "
                "procesamiento de 10 MB"
            )
        }

    content_type = response.get(
        "ContentType",
        "application/octet-stream"
    )

    data = response["Body"].read()

    extension = os.path.splitext(
        object_key
    )[1].lower()

    if (
        content_type == "text/plain"
        or extension == ".txt"
    ):
        text = extract_txt(data)

    elif (
        content_type == "application/pdf"
        or extension == ".pdf"
    ):
        text = extract_pdf(data)

    elif (
        content_type
        == (
            "application/vnd.openxmlformats-"
            "officedocument.wordprocessingml.document"
        )
        or extension == ".docx"
    ):
        text = extract_docx(data)

    else:
        return {
            "success": False,
            "error": "Formato de documento no soportado"
        }

    text = text.strip()

    if not text:
        return {
            "success": False,
            "error": (
                "No se encontró texto extraíble "
                "en el documento"
            )
        }

    return {
        "success": True,
        "action": "text_extracted",
        "bucket": bucket_name,
        "object_key": object_key,
        "content_type": content_type,
        "characters": len(text),
        "text": text
    }

def http_response(
    status_code: int,
    body: dict
) -> dict:
    return {
        "statusCode": status_code,
        "headers": {
            "Content-Type": "application/json"
        },
        "body": json.dumps(
            body,
            ensure_ascii=False
        )
    }

def lambda_handler(event, context):
    try:
        is_http_request = (
            isinstance(event, dict)
            and "requestContext" in event
            and "http" in event.get(
                "requestContext",
                {}
            )
        )

        if is_http_request:
            try:
                body = json.loads(
                    event.get("body") or "{}"
                )
            except json.JSONDecodeError:
                return http_response(
                    400,
                    {
                        "success": False,
                        "error": "JSON inválido"
                    }
                )

            result = generate_upload_url(
                body
            )

            status_code = (
                200
                if result.get("success")
                else 400
            )

            return http_response(
                status_code,
                result
            )

        action = event.get(
            "action",
            "generate_upload_url"
        )

        if action == "generate_upload_url":
            return generate_upload_url(
                event
            )

        if action == "extract_text":
            return extract_document_text(
                event
            )

        return {
            "success": False,
            "error": (
                f"Acción no soportada: {action}"
            )
        }

    except Exception as error:
        print(
            f"Documents function error: "
            f"{type(error).__name__}"
        )

        is_http_request = (
            isinstance(event, dict)
            and "requestContext" in event
        )

        error_body = {
            "success": False,
            "error": (
                "No se pudo completar la "
                "operación del documento"
            ),
            "error_type": (
                type(error).__name__
            )
        }

        if is_http_request:
            return http_response(
                500,
                error_body
            )

        return error_body

