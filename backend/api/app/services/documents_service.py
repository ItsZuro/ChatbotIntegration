import os
from datetime import (
    datetime,
    timezone,
)
from io import BytesIO
from uuid import uuid4

import boto3

from docx import Document
from pypdf import PdfReader

from app.core.config import (
    get_settings,
)


settings = get_settings()


s3_client = boto3.client(
    "s3",
    region_name=settings.aws_region,
)


ALLOWED_CONTENT_TYPES = {
    "application/pdf",
    (
        "application/vnd.openxmlformats-"
        "officedocument.wordprocessingml.document"
    ),
    "text/plain",
}


PRESIGNED_URL_EXPIRATION = 900

MAX_PROCESSING_FILE_SIZE = (
    10 * 1024 * 1024
)


def _get_bucket_name() -> str:
    bucket_name = (
        settings.documents_bucket_name
    )

    if not bucket_name:
        raise RuntimeError(
            "DOCUMENTS_BUCKET_NAME "
            "no está configurado."
        )

    return bucket_name


def _get_user_prefix(
    user_id: str,
) -> str:
    if not user_id:
        raise ValueError(
            "user_id es obligatorio."
        )

    return f"uploads/{user_id}/"


def _validate_object_key(
    object_key: str,
    user_id: str,
) -> None:
    if not object_key:
        raise ValueError(
            "object_key es obligatorio."
        )

    expected_prefix = (
        _get_user_prefix(
            user_id
        )
    )

    if not object_key.startswith(
        expected_prefix
    ):
        raise ValueError(
            "No tienes acceso "
            "a este documento."
        )


def _get_original_file_name(
    object_key: str,
) -> str:
    stored_name = os.path.basename(
        object_key
    )

    # UUID = 36 caracteres
    # Formato:
    # uuid-nombre-original.pdf
    if (
        len(stored_name) > 37
        and stored_name[36] == "-"
    ):
        return stored_name[37:]

    return stored_name


def generate_upload_url(
    user_id: str,
    file_name: str,
    content_type: str,
) -> dict:
    if not file_name:
        raise ValueError(
            "file_name es obligatorio."
        )

    if (
        content_type
        not in ALLOWED_CONTENT_TYPES
    ):
        raise ValueError(
            "Tipo de archivo no permitido."
        )

    bucket_name = (
        _get_bucket_name()
    )

    now = datetime.now(
        timezone.utc
    )

    safe_file_name = (
        os.path.basename(
            file_name
        )
    )

    object_key = (
        f"{_get_user_prefix(user_id)}"
        f"{now.year}/"
        f"{now.month:02d}/"
        f"{uuid4()}-"
        f"{safe_file_name}"
    )

    upload_url = (
        s3_client
        .generate_presigned_url(
            ClientMethod=(
                "put_object"
            ),
            Params={
                "Bucket":
                    bucket_name,
                "Key":
                    object_key,
                "ContentType":
                    content_type,
            },
            ExpiresIn=(
                PRESIGNED_URL_EXPIRATION
            ),
        )
    )

    return {
        "success": True,
        "action": (
            "upload_url_generated"
        ),
        "bucket": bucket_name,
        "object_key": object_key,
        "upload_url": upload_url,
        "expires_in": (
            PRESIGNED_URL_EXPIRATION
        ),
    }


def list_documents(
    user_id: str,
    limit: int = 100,
) -> dict:
    bucket_name = (
        _get_bucket_name()
    )

    safe_limit = max(
        1,
        min(
            int(limit),
            100,
        ),
    )

    response = (
        s3_client
        .list_objects_v2(
            Bucket=bucket_name,
            Prefix=_get_user_prefix(
                user_id
            ),
            MaxKeys=safe_limit,
        )
    )

    documents = []

    for item in response.get(
        "Contents",
        [],
    ):
        object_key = item[
            "Key"
        ]

        documents.append(
            {
                "object_key":
                    object_key,
                "file_name":
                    _get_original_file_name(
                        object_key
                    ),
                "size":
                    item.get(
                        "Size",
                        0,
                    ),
                "last_modified":
                    item[
                        "LastModified"
                    ].isoformat(),
            }
        )

    documents.sort(
        key=lambda item: (
            item[
                "last_modified"
            ]
        ),
        reverse=True,
    )

    return {
        "success": True,
        "documents": documents,
    }


def generate_download_url(
    user_id: str,
    object_key: str,
) -> dict:
    _validate_object_key(
        object_key=object_key,
        user_id=user_id,
    )

    bucket_name = (
        _get_bucket_name()
    )

    s3_client.head_object(
        Bucket=bucket_name,
        Key=object_key,
    )

    download_url = (
        s3_client
        .generate_presigned_url(
            ClientMethod=(
                "get_object"
            ),
            Params={
                "Bucket":
                    bucket_name,
                "Key":
                    object_key,
            },
            ExpiresIn=(
                PRESIGNED_URL_EXPIRATION
            ),
        )
    )

    return {
        "success": True,
        "action": (
            "download_url_generated"
        ),
        "object_key": (
            object_key
        ),
        "download_url": (
            download_url
        ),
        "expires_in": (
            PRESIGNED_URL_EXPIRATION
        ),
    }


def _extract_txt(
    data: bytes,
) -> str:
    try:
        return data.decode(
            "utf-8-sig"
        )

    except UnicodeDecodeError:
        return data.decode(
            "latin-1",
            errors="replace",
        )


def _extract_pdf(
    data: bytes,
) -> str:
    reader = PdfReader(
        BytesIO(data)
    )

    pages = []

    for page in reader.pages:
        text = (
            page.extract_text()
        )

        if text:
            pages.append(
                text.strip()
            )

    return "\n\n".join(
        pages
    )


def _extract_docx(
    data: bytes,
) -> str:
    document = Document(
        BytesIO(data)
    )

    blocks = []

    for paragraph in (
        document.paragraphs
    ):
        text = (
            paragraph
            .text
            .strip()
        )

        if text:
            blocks.append(
                text
            )

    for table in document.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell
                in row.cells
            ]

            row_text = (
                " | ".join(
                    value
                    for value
                    in values
                    if value
                )
            )

            if row_text:
                blocks.append(
                    row_text
                )

    return "\n".join(
        blocks
    )


def extract_document_text(
    user_id: str,
    object_key: str,
) -> str:
    _validate_object_key(
        object_key=object_key,
        user_id=user_id,
    )

    bucket_name = (
        _get_bucket_name()
    )

    response = (
        s3_client.get_object(
            Bucket=bucket_name,
            Key=object_key,
        )
    )

    content_length = (
        response.get(
            "ContentLength",
            0,
        )
    )

    if (
        content_length >
        MAX_PROCESSING_FILE_SIZE
    ):
        raise ValueError(
            "El archivo supera "
            "el límite de "
            "procesamiento de 10 MB."
        )

    content_type = (
        response.get(
            "ContentType",
            (
                "application/"
                "octet-stream"
            ),
        )
    )

    data = (
        response[
            "Body"
        ].read()
    )

    extension = (
        os.path.splitext(
            object_key
        )[1].lower()
    )

    if (
        content_type ==
        "text/plain"
        or extension == ".txt"
    ):
        text = _extract_txt(
            data
        )

    elif (
        content_type ==
        "application/pdf"
        or extension == ".pdf"
    ):
        text = _extract_pdf(
            data
        )

    elif (
        content_type ==
        (
            "application/vnd."
            "openxmlformats-"
            "officedocument."
            "wordprocessingml."
            "document"
        )
        or extension == ".docx"
    ):
        text = _extract_docx(
            data
        )

    else:
        raise ValueError(
            "Formato de documento "
            "no soportado."
        )

    text = text.strip()

    if not text:
        raise ValueError(
            "No se encontró texto "
            "extraíble en el documento."
        )

    return text


def delete_document(
    user_id: str,
    object_key: str,
) -> dict:
    _validate_object_key(
        object_key=object_key,
        user_id=user_id,
    )

    bucket_name = (
        _get_bucket_name()
    )

    objects_to_delete = []

    key_marker = None
    version_id_marker = None

    while True:
        params = {
            "Bucket":
                bucket_name,
            "Prefix":
                object_key,
        }

        if key_marker:
            params[
                "KeyMarker"
            ] = key_marker

        if version_id_marker:
            params[
                "VersionIdMarker"
            ] = version_id_marker

        response = (
            s3_client
            .list_object_versions(
                **params
            )
        )

        for version in (
            response.get(
                "Versions",
                [],
            )
        ):
            if (
                version.get(
                    "Key"
                )
                == object_key
            ):
                objects_to_delete.append(
                    {
                        "Key":
                            object_key,
                        "VersionId":
                            version[
                                "VersionId"
                            ],
                    }
                )

        for marker in (
            response.get(
                "DeleteMarkers",
                [],
            )
        ):
            if (
                marker.get(
                    "Key"
                )
                == object_key
            ):
                objects_to_delete.append(
                    {
                        "Key":
                            object_key,
                        "VersionId":
                            marker[
                                "VersionId"
                            ],
                    }
                )

        if not response.get(
            "IsTruncated"
        ):
            break

        key_marker = (
            response.get(
                "NextKeyMarker"
            )
        )

        version_id_marker = (
            response.get(
                "NextVersionIdMarker"
            )
        )

    if not objects_to_delete:
        raise ValueError(
            "El documento no existe."
        )

    deleted_count = 0

    # delete_objects admite
    # máximo 1000 objetos.
    for index in range(
        0,
        len(objects_to_delete),
        1000,
    ):
        batch = (
            objects_to_delete[
                index:index + 1000
            ]
        )

        response = (
            s3_client.delete_objects(
                Bucket=bucket_name,
                Delete={
                    "Objects":
                        batch,
                    "Quiet":
                        True,
                },
            )
        )

        errors = response.get(
            "Errors",
            [],
        )

        if errors:
            raise RuntimeError(
                "No se pudieron eliminar "
                "todas las versiones "
                "del documento."
            )

        deleted_count += len(
            batch
        )

    return {
        "success": True,
        "action": (
            "document_deleted"
        ),
        "object_key": (
            object_key
        ),
        "deleted_versions": (
            deleted_count
        ),
    }